import docx
import pymupdf as fitz

# 造一个测试 Word 文件
d = docx.Document()
d.add_heading("公司产品说明", level=1)
d.add_paragraph("我们的退货政策是：签收后 7 天内可以无理由退货。")
d.add_paragraph("客服电话是 400-123-4567，工作时间每天 9:00 到 21:00。")
d.save("test.docx")

# 造一个测试 PDF 文件
pdf = fitz.open()
page = pdf.new_page()
page.insert_text((72, 72), "这是一份测试 PDF 文档。", fontname="china-s")
page.insert_text((72, 100), "它包含两行文字，用来验证 PDF 解析。", fontname="china-s")
pdf.save("test.pdf")
pdf.close()

# 造一个 200 页大 PDF，用于验收大文件流式入库
big = fitz.open()
for i in range(200):
    p = big.new_page()
    p.insert_text((72, 72), f"大文件测试 第 {i+1} 页：", fontname="china-s")
    p.insert_text((72, 100), f"这是第 {i+1} 页的内容，用于验证流式解析与分批入库。", fontname="china-s")
    p.insert_text((72, 128), "每页都有几行中文文字，确保分块后能正确向量化。", fontname="china-s")
big.save("big.pdf")
big.close()

print("测试文件已生成：test.docx 和 test.pdf 和 big.pdf")