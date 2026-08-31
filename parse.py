# -*- coding: utf-8 -*-
"""文档解析：把 PDF / Word 转成干净的文本（支持流式逐段提取）。"""

import pymupdf as fitz
import docx


def _clean(text: str) -> str:
    """清洗：去掉行尾空格和多余的空行。"""
    lines = [line.rstrip() for line in text.split("\n")]
    lines = [line for line in lines if line.strip()]
    return "\n".join(lines)


def parse_pdf(path: str) -> str:
    """解析 PDF：逐页提取文字，并拼接。"""
    pdf = fitz.open(path)
    pages = [page.get_text() for page in pdf]
    pdf.close()
    return _clean("\n\n".join(pages))


def parse_docx(path: str) -> str:
    """解析 Word(.docx)：提取所有段落文字。"""
    d = docx.Document(path)
    paras = [p.text for p in d.paragraphs]
    return _clean("\n\n".join(paras))


def parse_file(path: str) -> str:
    """按文件后缀分发到对应的解析器。"""
    low = path.lower()
    if low.endswith(".pdf"):
        return parse_pdf(path)
    if low.endswith(".docx"):
        return parse_docx(path)
    raise ValueError(f"不支持的文件类型：{path}")


def parse_file_stream(path: str):
    """流式解析：逐页/逐段 yield 清洗后的文本，避免大文件整篇拼接。

    PDF：每页 yield 一段；Word：攒够一段长度再 yield。
    """
    low = path.lower()
    if low.endswith(".pdf"):
        pdf = fitz.open(path)
        for page in pdf:
            t = _clean(page.get_text())
            if t:
                yield t
        pdf.close()
    elif low.endswith(".docx"):
        d = docx.Document(path)
        buf = []
        for p in d.paragraphs:
            buf.append(p.text)
            if sum(len(x) for x in buf) >= 1000:
                seg = _clean("\n".join(buf))
                if seg:
                    yield seg
                buf = []
        if buf:
            seg = _clean("\n".join(buf))
            if seg:
                yield seg
    else:
        raise ValueError(f"不支持的文件类型：{path}")